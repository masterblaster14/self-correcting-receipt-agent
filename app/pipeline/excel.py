"""Excel output (openpyxl): one workbook per receipt, and a ledger-wide export."""
from __future__ import annotations

import io
import json
from datetime import datetime
from typing import List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import ProcessResult

HEAD_FILL = PatternFill("solid", fgColor="EEF2FF")
HEAD_FONT = Font(bold=True, color="3730A3")
OK_FILL = PatternFill("solid", fgColor="DCFCE7")
BAD_FILL = PatternFill("solid", fgColor="FEE2E2")
WARN_FILL = PatternFill("solid", fgColor="FEF3C7")
THIN = Side(style="thin", color="E2E8F0")
BORDER = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
MONEY = '#,##0.00'


def _header(ws, row: int, labels: List[str]):
    for i, label in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=label)
        c.fill, c.font, c.border = HEAD_FILL, HEAD_FONT, BORDER
        c.alignment = Alignment(vertical="center")


def _autosize(ws, min_w=8, max_w=60):
    widths = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                widths[c.column] = max(widths.get(c.column, 0), len(str(c.value)))
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(min_w, min(max_w, w + 2))


def build_receipt_xlsx(res: ProcessResult) -> bytes:
    p = res.profile
    wb = Workbook()

    # ---- Summary
    ws = wb.active
    ws.title = "Summary"
    rows = [
        ("Receipt ID", res.id), ("Status", res.state.value), ("Vendor", p.vendor_name), ("Address", p.vendor_address),
        ("Invoice #", p.invoice_number), ("Date", p.date), ("Time", p.time), ("Currency", p.currency),
        ("Payment method", p.payment_method), ("Category", res.category), ("Category reason", res.category_reason),
        ("Policy status", res.policy.status), ("Policy reason", res.policy.reason),
        ("Duplicate", "YES - " + res.duplicate.detail if res.duplicate.is_duplicate else "no"),
        ("Subtotal", p.subtotal), ("Tax total", p.tax_sum()), ("Other charges", p.charges_sum()),
        ("Discount", p.discount), ("Round off", p.round_off), ("TOTAL", p.total),
        ("Self-correction", "enabled" if res.self_correction_enabled else "disabled (baseline)"),
        ("Correction iterations", res.iterations), ("Model", res.model), ("Processing time (s)", round(res.total_ms / 1000, 1)),
        ("Generated", datetime.now().strftime("%Y-%m-%d %H:%M")),
    ]
    _header(ws, 1, ["Field", "Value"])
    for i, (k, v) in enumerate(rows, 2):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        c = ws.cell(row=i, column=2, value=v)
        if isinstance(v, (int, float)) and k not in ("Correction iterations",):
            c.number_format = MONEY
        if k == "TOTAL":
            c.font = Font(bold=True)
        if k == "Status":
            c.fill = OK_FILL if v == "VERIFIED" else WARN_FILL
    _autosize(ws)

    # ---- Line items
    ws = wb.create_sheet("Line items")
    _header(ws, 1, ["#", "Description", "Qty", "Unit price", "Amount"])
    for i, li in enumerate(p.line_items, 1):
        ws.append([i, li.description, li.quantity, li.unit_price, li.amount])
    n = len(p.line_items)
    r = n + 2
    ws.cell(row=r, column=4, value="Σ items").font = Font(bold=True)
    ws.cell(row=r, column=5, value=f"=SUM(E2:E{n + 1})" if n else 0).font = Font(bold=True)
    r += 1
    if p.subtotal is not None:
        ws.cell(row=r, column=4, value="Subtotal"); ws.cell(row=r, column=5, value=p.subtotal); r += 1
    for t in p.taxes:
        ws.cell(row=r, column=4, value=f"{t.label}" + (f" ({t.rate_percent:g}%)" if t.rate_percent is not None else "")); ws.cell(row=r, column=5, value=t.amount); r += 1
    for c in p.other_charges:
        ws.cell(row=r, column=4, value=c.label); ws.cell(row=r, column=5, value=c.amount); r += 1
    if p.discount:
        ws.cell(row=r, column=4, value="Discount"); ws.cell(row=r, column=5, value=-p.discount); r += 1
    if p.round_off:
        ws.cell(row=r, column=4, value="Round off"); ws.cell(row=r, column=5, value=p.round_off); r += 1
    ws.cell(row=r, column=4, value="TOTAL").font = Font(bold=True)
    ws.cell(row=r, column=5, value=p.total).font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=5):
        for c in row:
            c.number_format = MONEY
    _autosize(ws)

    # ---- Verification
    ws = wb.create_sheet("Verification")
    _header(ws, 1, ["ID", "Check", "Result", "Severity", "Message", "Hypothesis"])
    for c in res.verification.checks:
        ws.append([c.id, c.name, "PASS" if c.passed else "FAIL", c.severity.value, c.message, c.hypothesis])
        ws.cell(row=ws.max_row, column=3).fill = OK_FILL if c.passed else (BAD_FILL if c.severity.value == "error" else WARN_FILL)
    _autosize(ws)

    # ---- Trace
    ws = wb.create_sheet("Agent trace")
    _header(ws, 1, ["Iteration", "State", "Step", "Detail", "Failed checks", "Fields changed", "Agent notes", "Zoomed regions", "Seconds"])
    for e in res.trace:
        ws.append([e.iteration, e.state.value, e.title, e.detail, "; ".join(e.failed_checks),
                   "; ".join(f"{ch.field}: {ch.before} -> {ch.after}" for ch in e.changes),
                   e.revision_notes, ", ".join(e.focus_regions), round(e.duration_ms / 1000, 1) if e.duration_ms else None])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    _autosize(ws, max_w=50)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_ledger_xlsx(records: List[dict]) -> bytes:
    """records: rows from storage.get()-style dicts with a parsed 'result'."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Receipts"
    _header(ws, 1, ["Processed at", "Receipt ID", "Vendor", "Date", "Invoice #", "Category", "Currency", "Subtotal", "Tax",
                    "Charges", "Discount", "Total", "Status", "Policy", "Duplicate", "Correction passes", "Self-correction", "Model",
                    "Submitted by", "Department"])
    items_rows = []
    for r in records:
        res = r["result"]
        p = res["profile"]
        tax = sum((t.get("amount") or 0) for t in p.get("taxes", []))
        charges = sum((c.get("amount") or 0) for c in p.get("other_charges", []))
        ws.append([r["created_at"].replace("T", " "), r["id"], p.get("vendor_name"), p.get("date"), p.get("invoice_number"),
                   res.get("category"), p.get("currency"), p.get("subtotal"), tax, charges, p.get("discount"), p.get("total"),
                   res.get("state"), (res.get("policy") or {}).get("status"), "yes" if (res.get("duplicate") or {}).get("is_duplicate") else "no",
                   res.get("iterations"), "on" if res.get("self_correction_enabled") else "off (baseline)", res.get("model"),
                   r.get("submitted_by") or res.get("submitted_by"), r.get("department") or res.get("department")])
        st = ws.cell(row=ws.max_row, column=13)
        st.fill = OK_FILL if st.value == "VERIFIED" else WARN_FILL
        for i, li in enumerate(p.get("line_items", []), 1):
            items_rows.append([r["id"], p.get("vendor_name"), p.get("date"), res.get("category"), i, li.get("description"),
                               li.get("quantity"), li.get("unit_price"), li.get("amount")])
    n = len(records)
    if n:
        # totals per currency - never add INR to USD
        r = n + 2
        ws.cell(row=r, column=11, value="TOTAL by currency").font = Font(bold=True)
        for cur in sorted({(rec["result"]["profile"].get("currency") or "INR") for rec in records}):
            ws.cell(row=r, column=10, value=cur).font = Font(bold=True)
            ws.cell(row=r, column=12, value=f'=SUMIF(G2:G{n + 1},"{cur}",L2:L{n + 1})').font = Font(bold=True)
            r += 1
    for row in ws.iter_rows(min_row=2, min_col=8, max_col=12):
        for c in row:
            c.number_format = MONEY
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _autosize(ws, max_w=40)

    ws2 = wb.create_sheet("Line items")
    _header(ws2, 1, ["Receipt ID", "Vendor", "Date", "Category", "#", "Description", "Qty", "Unit price", "Amount"])
    for row in items_rows:
        ws2.append(row)
    for row in ws2.iter_rows(min_row=2, min_col=7, max_col=9):
        for c in row:
            c.number_format = MONEY
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = ws2.dimensions
    _autosize(ws2, max_w=50)

    # category pivot, per currency
    ws3 = wb.create_sheet("By category")
    _header(ws3, 1, ["Category", "Currency", "Receipts", "Total"])
    agg: dict = {}
    for r in records:
        cat = r["result"].get("category") or "Other"
        cur = r["result"]["profile"].get("currency") or "INR"
        a = agg.setdefault((cat, cur), [0, 0.0])
        a[0] += 1
        a[1] += r["result"]["profile"].get("total") or 0
    for (cat, cur), (cnt, tot) in sorted(agg.items(), key=lambda kv: (kv[0][1], -kv[1][1])):
        ws3.append([cat, cur, cnt, tot])
        ws3.cell(row=ws3.max_row, column=4).number_format = MONEY
    _autosize(ws3)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
